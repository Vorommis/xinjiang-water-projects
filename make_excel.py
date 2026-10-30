# -*- coding: utf-8 -*-
"""
把重建好的数据库导出成《新疆水利项目统一信息库》Excel（沿用原有分表结构）。
用法: python make_excel.py [--db xinjiang_water_projects.db] [--out 新疆水利项目统一信息库(更新版).xlsx]
"""
import argparse, os, sqlite3, sys

sys.stdout.reconfigure(encoding='utf-8')

# 与用户原有 Excel 完全一致的 42 列表头
HEADERS = [
    '序号', '地州', '项目类型', '分类号', '项目名称', '建设内容及规模', '总投资(万元)',
    '政府投资(万元)', '其他投资(万元)', '2026预计完成投资(万元)', '2026政府投资(万元)',
    '2026其他投资(万元)', '计划开工时间', '计划竣工时间', '建设地点', '项目法人',
    '项目单位性质', '联系人及联系方式', '资金来源', '设计单位', '招标代理',
    '项目进展情况', '关键关系资源', '关系资源级别', '关系资源备注',
    '业主决策层联系人', '业主决策层职务', '业主参与层联系人', '业主参与层职务',
    '业主操作层联系人', '业主操作层职务', '公共资源中心联系人', '主要竞争对手',
    '评标专家', '投标策略要点', '策划责任人', '市场经理', '攻关状态', '下一步行动',
    '预计招标时间', '中标概率', '备注',
]
# 数据库字段 -> Excel 列名
FIELDS = [
    ('', '序号'), ('地州', '地州'), ('项目类型', '项目类型'), ('分类号', '分类号'),
    ('项目名称', '项目名称'), ('建设内容及规模', '建设内容及规模'),
    ('总投资数值', '总投资(万元)'), ('政府投资_万元', '政府投资(万元)'),
    ('其他投资_万元', '其他投资(万元)'),
    ('年2026预计完成投资_万元', '2026预计完成投资(万元)'),
    ('年2026政府投资_万元', '2026政府投资(万元)'),
    ('年2026其他投资_万元', '2026其他投资(万元)'),
    ('计划开工时间', '计划开工时间'), ('计划竣工时间', '计划竣工时间'),
    ('建设地点', '建设地点'), ('项目法人', '项目法人'),
    ('项目单位性质', '项目单位性质'), ('联系人及联系方式', '联系人及联系方式'),
    ('资金来源', '资金来源'), ('设计单位', '设计单位'), ('招标代理', '招标代理'),
    ('项目进展情况', '项目进展情况'), ('关键关系资源', '关键关系资源'),
    ('关系资源级别', '关系资源级别'), ('关系资源备注', '关系资源备注'),
    ('业主决策层联系人', '业主决策层联系人'), ('业主决策层职务', '业主决策层职务'),
    ('业主参与层联系人', '业主参与层联系人'), ('业主参与层职务', '业主参与层职务'),
    ('业主操作层联系人', '业主操作层联系人'), ('业主操作层职务', '业主操作层职务'),
    ('公共资源中心联系人', '公共资源中心联系人'), ('主要竞争对手', '主要竞争对手'),
    ('评标专家', '评标专家'), ('投标策略要点', '投标策略要点'),
    ('策划责任人', '策划责任人'), ('市场经理', '市场经理'), ('攻关状态', '攻关状态'),
    ('下一步行动', '下一步行动'), ('预计招标时间', '预计招标时间'),
    ('中标概率(%)', '中标概率'), ('备注', '备注'),
]
EXTRA_TAIL = ['数据来源', '数据质量', '重点跟踪']

REGIONS = ['乌鲁木齐市', '昌吉州', '阿克苏地区', '阿勒泰地区', '巴州', '博州', '哈密市',
           '和田地区', '喀什地区', '克拉玛依市', '克州', '塔城地区', '吐鲁番市', '伊犁州']
NUM_COLS = {'总投资(万元)', '政府投资(万元)', '其他投资(万元)',
            '2026预计完成投资(万元)', '2026政府投资(万元)', '2026其他投资(万元)'}
WIDTHS = {'序号': 6, '地州': 10, '项目类型': 8, '分类号': 12, '项目名称': 42,
          '建设内容及规模': 46, '总投资(万元)': 13, '政府投资(万元)': 14,
          '其他投资(万元)': 13, '2026预计完成投资(万元)': 18,
          '2026政府投资(万元)': 16, '2026其他投资(万元)': 16,
          '计划开工时间': 12, '计划竣工时间': 12, '建设地点': 14, '项目法人': 26,
          '项目单位性质': 12, '联系人及联系方式': 22, '资金来源': 16, '设计单位': 18,
          '招标代理': 18, '项目进展情况': 30, '备注': 18, '数据来源': 20, '数据质量': 34, '重点跟踪': 10}


def num_or_none(v):
    if v is None or str(v).strip() in ('', '-'):
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (ValueError, TypeError):
        return v


def fetch(db):
    con = sqlite3.connect(db)
    con.row_factory = sqlite3.Row
    rows = [dict(r) for r in con.execute('SELECT * FROM projects ORDER BY row_id')]
    con.close()
    return rows


def region_of(rec):
    dz = (rec.get('地州') or '').strip()
    return dz if dz in REGIONS else '兵团'


def build_sheet(ws, recs, title=None, renumber=False):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    head = HEADERS + EXTRA_TAIL
    ws.append(head)
    fill = PatternFill('solid', fgColor='1A5276')
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, rec in enumerate(recs, 1):
        line = []
        for db_f, _ in FIELDS:
            if db_f == '':
                line.append(i if renumber else (rec.get('原始序号') or ''))
            elif db_f in NUM_COLS or db_f == '总投资数值':
                line.append(num_or_none(rec.get(db_f)))
            else:
                v = rec.get(db_f)
                line.append('' if v in (None, '-') else v)
        line.append(rec.get('数据来源', ''))
        line.append(rec.get('数据质量', ''))
        line.append(rec.get('重点跟踪', ''))
        ws.append(line)
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(head))}{ws.max_row}'
    for i, name in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, 14)
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--db', default='xinjiang_water_projects.db')
    ap.add_argument('--out', default='新疆水利项目统一信息库(更新版).xlsx')
    a = ap.parse_args()

    import openpyxl
    rows = fetch(a.db)
    cur = [r for r in rows if r['当前看板'] == 1]
    old = [r for r in rows if r['当前看板'] == 0]
    print(f'读取 {len(rows)} 条：当前 {len(cur)}，已筛除 {len(old)}')

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1) 分地州表（当前记录），与用户原结构一致
    for rg in REGIONS + ['兵团']:
        sub = [r for r in cur if region_of(r) == rg]
        ws = wb.create_sheet(rg)
        build_sheet(ws, sub, renumber=True)
        print(f'  {rg:8} {len(sub):3} 条')

    # 2) 总表（当前）
    build_sheet(wb.create_sheet('水利项目总表'), cur)
    # 3) 已筛除记录
    build_sheet(wb.create_sheet('已筛除记录'), old)
    # 4) 全部记录
    build_sheet(wb.create_sheet('全部记录'), rows)

    wb.save(a.out)
    print(f'\n已写出 {a.out}  ({os.path.getsize(a.out):,} bytes)')
    print('工作表:', wb.sheetnames)


if __name__ == '__main__':
    main()

